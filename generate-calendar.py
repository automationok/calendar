from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import openai
import os
from docx import Document
import tempfile
import uuid

app = Flask(__name__)
CORS(app)

openai.api_key = os.getenv("OPENAI_API_KEY")

@app.route('/generate-calendar', methods=['POST'])
def generate_calendar():
    data = request.json
    business_name = data.get("business_name", "Your Business")
    industry = data.get("industry", "general")
    subscription = data.get("subscription", "Free")

    prompt = f"""
    You are an expert social media strategist.

    Generate a 1-week AI-powered content calendar for the business below:
    - Business Name: {business_name}
    - Industry: {industry}
    - Subscription Plan: {subscription}

    🎯 Strategy Guidelines:
    - Every post should address specific audience pain points.
    - Follow the latest social media engagement trends.
    - Mix short-form and long-form content.
    - Include **video content** suggestions in each plan.
    - Optimize for engagement (best times, emotional hooks, hashtags).
    
    📦 If the subscription is "Premium", include exactly:
    - 📌 Monday: AI-Generated Customer Testimonial Video
    - 📌 Wednesday: Live Q&A Session (AI Recommends Topics)
    - 📌 Friday: AI-Optimized Paid Ad Campaign (Social & Google)
    - 📌 Sunday: Personal Branding Blog Post (SEO-Optimized)
    
    🗓 Format your output like this:
    Day of Week: [Post Type] - [Short Description]
    Example:
    📌 Monday: Story-based video post – Highlight a transformation journey of a client.
    
    Generate the calendar now:
    """

    try:
        response = openai.chat.completions.create(
            model="gpt-4",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
            max_tokens=400
        )

        calendar_text = response.choices[0].message.content.strip()

        # Cria arquivo temporário .docx
        unique_id = str(uuid.uuid4())
        temp_dir = tempfile.gettempdir()
        filename = os.path.join(temp_dir, f"calendar_{unique_id}.docx")

        doc = Document()
        doc.add_heading(f"Content Calendar for {business_name}", 0)
        for line in calendar_text.split("\n"):
            doc.add_paragraph(line)

        doc.save(filename)

        # Retorna o arquivo como resposta
        return send_file(filename, as_attachment=True)

    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
